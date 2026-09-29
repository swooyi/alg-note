"""Import the owner's CSP sheet and SVGs; rerun with a new XLSX export to refresh.

python scripts/import_sq1_csp.py sheet.xlsx "path/to/All shape"
Requires openpyxl only for reading the source workbook.
"""
import argparse
import base64
import copy
import json
import shutil
import xml.etree.ElementTree as ET
from pathlib import Path
from xml.sax.saxutils import escape

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "alg/sq1/csp"
SOURCE = "https://docs.google.com/spreadsheets/d/1uP4QGcOrWnubi4ECC338XrfC9KXEWn2bfFNQ3MWk7Ho/edit?gid=1494475855#gid=1494475855"
NS = "http://www.w3.org/2000/svg"
ET.register_namespace("", NS)


def compact(value):
    return " ".join(str(value or "").split())


def shape_info(name):
    first = name.split()[0]
    side = "left" if "Left" in name else "right" if "Right" in name else ""
    if first in ("Fist", "Paw"):
        return f"{side}-{first.lower()}", f"{side} {first.lower()}.svg", f"{side.title()} {first}"
    if first in ("4-2", "5-1"):
        return f"{first}-{side}", f"[3C 6e] {first} {side}.svg", f"{first} {side.title()}"
    specials = {
        "Paired": ("pair", "pair.svg", "Pair"),
        "Parallel": ("i", "I.svg", "I"),
        "Perpendicular": ("l", "L.svg", "L"),
        "Star": ("star", "[6C] star.svg", "Star"),
        "1/7": ("7-1", "7-1.svg", "1/7"),
        "3/5": ("5-3", "[2C 8e] 5-3.svg", "3/5"),
        "2-6": ("6-2", "[2C 8e] 6-2.svg", "2-6"),
        "4-4": ("4-4", "[2C 8e] 4-4.svg", "4-4"),
        "8": ("8", "[2C 8e] 8.svg", "8"),
    }
    if first in specials:
        return specials[first]
    if first[0].isdigit():
        return first, f"[3C 6e] {first}.svg", first
    return first.lower(), f"{first.lower()}.svg", first


def display_svg(path, shape_id):
    root = ET.fromstring(path.read_text(encoding="utf-8"))
    # The two Illustrator exports have cropped pages. Restore the common
    # 200-unit frame around their piece origin; never rotate/mirror the art.
    origins = {"7-1": (50.76, 78.25), "star": (68.98, 78.25)}
    if shape_id in origins:
        x, y = origins[shape_id]
        root.set("viewBox", f"{x - 100} {y - 100} 200 200")
    root.set("width", "200")
    root.set("height", "200")
    # Isolated image documents keep Illustrator's .cls-* rules from leaking.
    return "data:image/svg+xml;base64," + base64.b64encode(ET.tostring(root)).decode("ascii")


def pair_svg(u, d, images):
    parts = [f'<svg xmlns="{NS}" viewBox="0 0 400 244" role="img" aria-label="{escape(u["name"] + " / " + d["name"], {chr(34): "&quot;"})}">']
    for x, label, shape in ((0, "U", u), (200, "D", d)):
        parts.append(f'<image x="{x}" y="20" width="200" height="200" href="{images[shape["id"]]}"/>')
        parts.append(f'<text x="{x + 100}" y="19" text-anchor="middle" font-family="sans-serif" font-size="18" fill="#24292f">{label}</text>')
        parts.append(f'<text x="{x + 100}" y="238" text-anchor="middle" font-family="sans-serif" font-size="17" fill="#24292f">{escape(shape["label"])}</text>')
    return "".join(parts) + "</svg>"


def write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("shapes", type=Path)
    args = parser.parse_args()
    workbook = openpyxl.load_workbook(args.workbook, data_only=True)
    sheet = next(s for s in workbook if s.title.strip() == "내CSP 학습")
    assert sheet["I3"].value == "Odd Alg" and sheet["J3"].value == "Even Alg"
    records, shapes, images, groups, cases, svgs = [], {}, {}, {}, [], {}
    for row in range(4, sheet.max_row + 1):
        cells = [sheet.cell(row, col) for col in range(1, 16)]
        values = [cell.value for cell in cells]
        if not values[2] or not values[3]:
            continue
        yellow = [cell.fill.fgColor.type == "rgb" and cell.fill.fgColor.rgb[-6:].upper() == "FFFF00" for cell in cells[8:10]]
        assert yellow[0] == yellow[1], f"Row {row}: common counting disagrees between Odd/Even"
        assert values[8] and values[9] and isinstance(values[5], (int, float)), row
        records.append({"row": row, "values": values, "extraCount": int(yellow[0])})
        pair = []
        for name in values[2:4]:
            name = compact(name)
            shape_id, filename, label = shape_info(name)
            if shape_id not in shapes:
                shapes[shape_id] = {"id": shape_id, "name": name, "label": label, "file": filename}
                images[shape_id] = display_svg(args.shapes / filename, shape_id)
            pair.append(copy.deepcopy(shapes[shape_id]))
        u, d = pair
        group_name = compact(values[4])
        if group_name not in groups:
            groups[group_name] = {"id": f"set-{len(groups) + 1}", "name": group_name, "caseIds": []}
        case_id = str(len(cases) + 1)
        groups[group_name]["caseIds"].append(case_id)
        cases.append({
            "id": case_id, "name": f'{u["label"]} / {d["label"]}',
            "group": groups[group_name]["id"],
            "algorithms": [str(values[8]).strip(), str(values[9]).strip()],
            "scramble": "", "scrambles": [], "svgId": case_id,
            "tags": {"recognition": list(dict.fromkeys([u["label"], d["label"]]))},
            "csp": {"u": u, "d": d, "probability": values[5],
                    "probabilityLabel": f"{values[5] * 100:.3f}%",
                    "extraCount": int(yellow[0]), "notes": str(values[10] or "").strip(),
                    "sourceRow": row},
        })
        svgs[case_id] = pair_svg(u, d, images)
    assert len(cases) == 90 and len(shapes) == 29
    assert len({tuple(sorted([c["csp"]["u"]["id"], c["csp"]["d"]["id"]])) for c in cases}) == 90
    assert sum(c["csp"]["extraCount"] for c in cases) == 28
    # Nothing is written until all source rows and image mappings pass.
    for shape in shapes.values():
        target = OUT / "shapes" / shape["file"]
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(args.shapes / shape["file"], target)
    common = {"schemaVersion": 1, "puzzle": "SQ1", "algset": "csp"}
    write_json(OUT / "algset.json", {
        "schemaVersion": 1, "puzzle": "SQ1", "id": "csp", "name": "CSP",
        "source": {"url": SOURCE, "sheet": sheet.title.strip(), "range": "C4:K93"},
        "notes": ["사용자 제공 시트 및 직접 제작한 모양 SVG.",
                  "algorithms[0]은 Odd, algorithms[1]은 Even 설명. csp.extraCount는 케이스 공통 카운팅 보정.",
                  "셋업의 패리티는 미확정이며 드릴은 추후 지원. 원본 셋업은 수집 자료에만 보관."],
    })
    write_json(OUT / "groups.json", {**common, "groups": list(groups.values())})
    write_json(OUT / "cases.json", {**common, "cases": cases})
    write_json(OUT / "svgs.json", {**common, "svgs": svgs})
    write_json(ROOT / "source/cache/sq1-csp/sheet.json", {
        "url": SOURCE, "sheet": sheet.title.strip(), "headers": [sheet.cell(3, c).value for c in range(1, 16)],
        "countingLegend": sheet["I2"].value, "records": records,
    })
    print(f"Imported {len(cases)} cases, {len(shapes)} shapes, 28 common counting corrections")


if __name__ == "__main__":
    main()
